#!/usr/bin/env python3
"""
render_pdf.py - branded Markdown -> PDF (and optional DOCX) for ledger documents.

    python3 render_pdf.py --markdown doc.md --out doc.pdf [--docx doc.docx] [--png doc.png]
                          [--logo assets/logo.png] [--title "..."] [--eyebrow "COMPLETETECH LLC"]
                          [--doc-type "USAGE STATEMENT"] [--footer "..."] [--accent "#1E3A8A"]

PDF needs reportlab; DOCX needs python-docx; the PNG preview needs pypdfium2 + pillow. Each is optional:
a missing library skips that output with a message instead of failing. The Markdown subset handled is what
the catalog templates use: #/##/### headings, paragraphs, bullet lists, pipe tables, **bold**, `code`.
"""
import argparse
import html
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import safety  # noqa: E402

BRAND = {"accent": "#1E3A8A", "ink": "#0F172A", "soft": "#EEF2FF", "muted": "#64748B", "border": "#E2E8F0", "zebra": "#F8FAFC",
         "name": "CompleteTech", "eyebrow": "COMPLETETECH LLC", "contact": "complete.tech · Timothy.Gregg@complete.tech"}


# --------------------------------------------------------------------------- Markdown -> blocks
def parse_blocks(md):
    """Yield ('h', level, text) | ('p', text) | ('ul', [items]) | ('table', rows) | ('quote', text)."""
    lines = md.replace("\r\n", "\n").split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s:
            i += 1
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            yield ("h", len(m.group(1)), m.group(2).strip())
            i += 1
            continue
        if s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c or "--") for c in cells):
                    rows.append(cells)
                i += 1
            yield ("table", rows)
            continue
        if s.startswith("- ") or s.startswith("* "):
            items = []
            while i < len(lines) and (lines[i].strip().startswith("- ") or lines[i].strip().startswith("* ")):
                items.append(lines[i].strip()[2:].strip())
                i += 1
            yield ("ul", items)
            continue
        if s.startswith(">"):
            yield ("quote", s.lstrip("> ").strip())
            i += 1
            continue
        buf = [s]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#{1,4}\s|\||- |\* |>)", lines[i].strip()):
            buf.append(lines[i].strip())
            i += 1
        yield ("p", " ".join(buf))


def inline(text):
    """Markdown inline -> reportlab / HTML markup (escaped)."""
    t = html.escape(text, quote=False)
    t = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", t)
    t = re.sub(r"`([^`]+)`", r"<font face='Courier'>\1</font>", t)
    return t


# --------------------------------------------------------------------------- PDF
def build_pdf(md, cfg, out_path):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from reportlab.lib.utils import ImageReader

    accent = colors.HexColor(cfg.get("accent") or BRAND["accent"])
    ink = colors.HexColor(BRAND["ink"])
    muted = colors.HexColor(BRAND["muted"])
    border = colors.HexColor(BRAND["border"])
    zebra = colors.HexColor(BRAND["zebra"])
    soft = colors.HexColor(BRAND["soft"])
    W, H = LETTER
    margin = 0.85 * inch
    avail = W - 2 * margin
    st = {
        "h1": ParagraphStyle("h1", fontName="Helvetica-Bold", fontSize=20, leading=24, textColor=ink, spaceAfter=10, spaceBefore=4),
        "h2": ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=13, leading=16, textColor=accent, spaceBefore=12, spaceAfter=6),
        "h3": ParagraphStyle("h3", fontName="Helvetica-Bold", fontSize=11, leading=14, textColor=ink, spaceBefore=8, spaceAfter=4),
        "p": ParagraphStyle("p", fontName="Helvetica", fontSize=9.6, leading=13.5, textColor=ink, spaceAfter=6),
        "li": ParagraphStyle("li", fontName="Helvetica", fontSize=9.6, leading=13.5, textColor=ink),
        "cell": ParagraphStyle("cell", fontName="Helvetica", fontSize=8.6, leading=11, textColor=ink),
        "cellh": ParagraphStyle("cellh", fontName="Helvetica-Bold", fontSize=8.6, leading=11, textColor=colors.white),
        "quote": ParagraphStyle("quote", fontName="Helvetica-Oblique", fontSize=9.6, leading=13.5, textColor=muted, leftIndent=12, spaceAfter=6),
    }
    story = []
    for blk in parse_blocks(md):
        kind = blk[0]
        if kind == "h":
            level, text = blk[1], blk[2]
            story.append(Paragraph(inline(text), st["h1" if level == 1 else "h2" if level == 2 else "h3"]))
        elif kind == "p":
            story.append(Paragraph(inline(blk[1]), st["p"]))
        elif kind == "quote":
            story.append(Paragraph(inline(blk[1]), st["quote"]))
        elif kind == "ul":
            story.append(ListFlowable([ListItem(Paragraph(inline(x), st["li"]), leftIndent=12) for x in blk[1]], bulletType="bullet", start="•", leftIndent=14, bulletFontSize=8))
            story.append(Spacer(1, 6))
        elif kind == "table":
            rows = blk[1]
            if not rows:
                continue
            ncol = max(len(r) for r in rows)
            rows = [r + [""] * (ncol - len(r)) for r in rows]
            numeric = [all(re.fullmatch(r"[\s$€£%+\-.,0-9×x]*", r[c]) for r in rows[1:]) for c in range(ncol)]
            data = [[Paragraph(inline(c), st["cellh"]) for c in rows[0]]]
            for r in rows[1:]:
                data.append([Paragraph(inline(c), ParagraphStyle("cr", parent=st["cell"], alignment=2 if numeric[ci] else 0)) for ci, c in enumerate(r)])
            widths = [max(len(re.sub(r"\*\*|`", "", r[c])) for r in rows) for c in range(ncol)]
            tot = float(sum(widths)) or 1.0
            colw = [max(0.7 * inch, avail * (w / tot)) for w in widths]
            scale = avail / sum(colw)
            colw = [w * scale for w in colw]
            t = Table(data, colWidths=colw, repeatRows=1, hAlign="LEFT")
            style = [("BACKGROUND", (0, 0), (-1, 0), accent), ("LINEBELOW", (0, 0), (-1, -1), 0.4, border), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                     ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4), ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5)]
            for ri in range(1, len(data)):
                if ri % 2 == 0:
                    style.append(("BACKGROUND", (0, ri), (-1, ri), zebra))
            t.setStyle(TableStyle(style))
            story.append(t)
            story.append(Spacer(1, 8))

    logo = cfg.get("logo")
    logo_reader = None
    if logo:
        try:
            logo_reader = ImageReader(logo)
        except Exception:
            logo_reader = None

    def page(c, doc):
        c.saveState()
        # letterhead band
        c.setFillColor(soft)
        c.rect(0, H - 0.95 * inch, W, 0.95 * inch, stroke=0, fill=1)
        c.setFillColor(accent)
        c.rect(0, H - 0.95 * inch, W, 3, stroke=0, fill=1)
        x = margin
        if logo_reader:
            iw, ih = logo_reader.getSize()
            h = 0.5 * inch
            w = h * iw / float(ih)
            c.drawImage(logo_reader, x, H - 0.72 * inch, width=w, height=h, mask="auto")
            x += w + 10
        c.setFillColor(muted)
        c.setFont("Helvetica-Bold", 7.5)
        c.drawString(x, H - 0.42 * inch, (cfg.get("eyebrow") or BRAND["eyebrow"]).upper())
        c.setFillColor(ink)
        c.setFont("Helvetica-Bold", 11)
        c.drawString(x, H - 0.60 * inch, cfg.get("title") or "")
        c.setFillColor(accent)
        c.setFont("Helvetica-Bold", 7.5)
        c.drawRightString(W - margin, H - 0.42 * inch, (cfg.get("doc_type") or "LEDGER DOCUMENT").upper())
        c.setFillColor(muted)
        c.setFont("Helvetica", 7.5)
        c.drawRightString(W - margin, H - 0.60 * inch, cfg.get("date") or "")
        # footer
        c.setStrokeColor(border)
        c.line(margin, 0.7 * inch, W - margin, 0.7 * inch)
        c.setFont("Helvetica", 7.2)
        c.drawString(margin, 0.52 * inch, cfg.get("footer") or "")
        c.drawRightString(W - margin, 0.52 * inch, "Page %d" % doc.page)
        c.restoreState()

    doc = SimpleDocTemplate(str(out_path), pagesize=LETTER, leftMargin=margin, rightMargin=margin, topMargin=1.15 * inch, bottomMargin=0.95 * inch,
                            title=cfg.get("title") or "", author=cfg.get("name") or BRAND["name"], subject=cfg.get("doc_type") or "")
    doc.build(story, onFirstPage=page, onLaterPages=page)


# --------------------------------------------------------------------------- DOCX
def build_docx(md, cfg, out_path):
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    accent = (cfg.get("accent") or BRAND["accent"]).lstrip("#")
    rgb = RGBColor(int(accent[0:2], 16), int(accent[2:4], 16), int(accent[4:6], 16))
    d = docx.Document()
    for s in d.sections:
        s.left_margin = s.right_margin = Inches(0.9)
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        hp = s.header.paragraphs[0]
        if cfg.get("logo"):
            try:
                hp.add_run().add_picture(cfg["logo"], height=Inches(0.45))
                hp.add_run("   ")
            except Exception:
                pass
        r = hp.add_run((cfg.get("eyebrow") or BRAND["eyebrow"]).upper() + "   ·   " + (cfg.get("doc_type") or "LEDGER DOCUMENT").upper())
        r.font.size = Pt(8)
        r.font.bold = True
        r.font.color.rgb = rgb
        fp = s.footer.paragraphs[0]
        fr = fp.add_run(cfg.get("footer") or "")
        fr.font.size = Pt(8)
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def add_runs(par, text):
        pos = 0
        for m in re.finditer(r"\*\*(.+?)\*\*|`([^`]+)`", text):
            if m.start() > pos:
                par.add_run(text[pos:m.start()])
            if m.group(1):
                par.add_run(m.group(1)).bold = True
            else:
                rr = par.add_run(m.group(2))
                rr.font.name = "Consolas"
            pos = m.end()
        if pos < len(text):
            par.add_run(text[pos:])

    for blk in parse_blocks(md):
        kind = blk[0]
        if kind == "h":
            h = d.add_heading("", level=min(blk[1], 3))
            add_runs(h, blk[2])
            for r in h.runs:
                r.font.color.rgb = rgb if blk[1] > 1 else RGBColor(0x0F, 0x17, 0x2A)
        elif kind == "p":
            add_runs(d.add_paragraph(), blk[1])
        elif kind == "quote":
            p = d.add_paragraph()
            add_runs(p, blk[1])
            for r in p.runs:
                r.italic = True
        elif kind == "ul":
            for item in blk[1]:
                add_runs(d.add_paragraph(style="List Bullet"), item)
        elif kind == "table":
            rows = blk[1]
            if not rows:
                continue
            ncol = max(len(r) for r in rows)
            t = d.add_table(rows=len(rows), cols=ncol)
            t.style = "Light Grid Accent 1"
            for ri, r in enumerate(rows):
                for ci in range(ncol):
                    cell = t.cell(ri, ci)
                    cell.text = ""
                    add_runs(cell.paragraphs[0], r[ci] if ci < len(r) else "")
                    if ri == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
            d.add_paragraph()
    d.core_properties.title = cfg.get("title") or ""
    d.core_properties.author = cfg.get("name") or BRAND["name"]
    d.save(str(out_path))


def montage(pdf_path, png_path, scale=1.6, max_pages=2):
    import pypdfium2 as pdfium
    from PIL import Image
    pdf = pdfium.PdfDocument(str(pdf_path))
    pages = [pdf[i].render(scale=scale).to_pil() for i in range(min(len(pdf), max_pages))]
    w = sum(p.width for p in pages) + 16 * (len(pages) - 1)
    h = max(p.height for p in pages)
    out = Image.new("RGB", (w, h), (226, 232, 240))
    x = 0
    for p in pages:
        out.paste(p, (x, 0))
        x += p.width + 16
    out.save(str(png_path))


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--markdown", required=True)
    ap.add_argument("--out", help="PDF path")
    ap.add_argument("--docx", help="DOCX path")
    ap.add_argument("--png", help="PNG preview of the first pages (needs pypdfium2 + pillow)")
    ap.add_argument("--logo")
    ap.add_argument("--title", default="")
    ap.add_argument("--eyebrow", default=BRAND["eyebrow"])
    ap.add_argument("--doc-type", default="LEDGER DOCUMENT")
    ap.add_argument("--footer", default="%s · %s · API-equivalent figures at list price, not an invoice" % (BRAND["name"], BRAND["contact"]))
    ap.add_argument("--accent", default=BRAND["accent"])
    ap.add_argument("--date", default="")
    a = ap.parse_args()
    with open(a.markdown, encoding="utf-8") as fh:
        md = fh.read()
    cfg = {"logo": a.logo, "title": a.title, "eyebrow": a.eyebrow, "doc_type": a.doc_type, "footer": a.footer, "accent": a.accent, "date": a.date}
    for _out in (a.out, a.docx, a.png):  # rendered documents carry the ledger contents: keep them owner-only
        if _out:
            safety.private_dir(os.path.dirname(os.path.abspath(_out)) or ".")
    rc = 0
    if a.out:
        try:
            build_pdf(md, cfg, a.out)
            print("PDF:", a.out)
        except ImportError as ex:
            print("[skip PDF] %s; pip install reportlab" % ex, file=sys.stderr)
            rc = 3
    if a.docx:
        try:
            build_docx(md, cfg, a.docx)
            print("DOCX:", a.docx)
        except ImportError as ex:
            print("[skip DOCX] %s; pip install python-docx" % ex, file=sys.stderr)
            rc = 3
    if a.png and a.out:
        try:
            montage(a.out, a.png)
            print("PNG:", a.png)
        except ImportError as ex:
            print("[skip PNG] %s; pip install pypdfium2 pillow" % ex, file=sys.stderr)
    for _out in (a.out, a.docx, a.png):
        if _out:
            safety.private_file(_out)
    return rc


if __name__ == "__main__":
    sys.exit(main())
